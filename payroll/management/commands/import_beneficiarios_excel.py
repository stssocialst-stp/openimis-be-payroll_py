"""
Management command: import_beneficiarios_excel

Importa a lista de beneficiários do Excel para o openIMIS, executando o fluxo
completo de criação de dados:

  Fase 1 (sempre):
    - Criar ou actualizar Individual (nome, data_n, sexo, BI, NIB, PAN, etc.)
    - Criar ou actualizar Group (agregado_id como code)
    - Criar ou actualizar GroupIndividual (HEAD + PRIMARY)

  Fase 2 (--benefit-plan-id):
    - Criar Beneficiary (ligação Individual→BenefitPlan, status ACTIVE)

  Fase 3 (--payroll-id):
    - Criar BenefitConsumption (amount=valor, status=ACCEPTED)
    - Criar PayrollBenefitConsumption (ligar ao Payroll)

Uso:
    # Só indivíduos e grupos (dry-run)
    python manage.py import_beneficiarios_excel /path/lista.xlsx --dry-run

    # Fluxo completo
    python manage.py import_beneficiarios_excel /path/lista.xlsx \\
        --benefit-plan-id <uuid> \\
        --payroll-id <uuid> \\
        --username admin

Colunas esperadas no Excel:
    nome_chefe, distrito, categoria_eligibilidade, agregado_id,
    sexo_chefe, data_n, bilhete_didentidade, pan, subconta, NIBs, valor, Situação, INDEX CSU2026
"""
import logging
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

logger = logging.getLogger(__name__)

DISTRICT_ALIASES = {
    'ÁGUA GRANDE': ['AGUA GRANDE', 'ÁGUA GRANDE'],
    'MÉZOCHI': ['MEZOCHI', 'MÉZOCHI'],
    'CANTAGALO': ['CANTAGALO'],
    'LEMBÁ': ['LEMBA', 'LEMBÁ'],
    'LOBATA': ['LOBATA'],
    'CAUÉ': ['CAUE', 'CAUÉ'],
    'PAGUÉ': ['PAGUE', 'PAGUÉ'],
}


class Command(BaseCommand):
    help = 'Importa beneficiários do Excel: Individual → Group → GroupIndividual → [Beneficiary] → [BenefitConsumption]'

    def add_arguments(self, parser):
        parser.add_argument('excel_path', type=str,
                            help='Caminho para o ficheiro Excel (.xlsx)')
        parser.add_argument('--dry-run', action='store_true', default=False,
                            help='Simula sem guardar nada na base de dados')
        parser.add_argument('--username', type=str, default='admin',
                            help='Username openIMIS para auditoria (default: admin)')
        parser.add_argument('--sheet', type=str, default=None,
                            help='Nome da sheet (default: primeira sheet)')
        parser.add_argument('--benefit-plan-id', type=str, default=None,
                            help='UUID do BenefitPlan — activa Fase 2 (Beneficiary)')
        parser.add_argument('--payroll-id', type=str, default=None,
                            help='UUID do Payroll — activa Fase 3 (BenefitConsumption)')
        parser.add_argument('--benefit-type', type=str, default='CASH',
                            help='Tipo de benefício para BenefitConsumption (default: CASH)')

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl não instalado. Execute: pip install openpyxl")

        from individual.models import Individual, Group, GroupIndividual
        from location.models import Location
        from payroll.models import BenefitConsumption, BenefitConsumptionStatus, PayrollBenefitConsumption, Payroll

        excel_path = Path(options['excel_path'])
        if not excel_path.exists():
            raise CommandError(f"Ficheiro não encontrado: {excel_path}")

        dry_run = options['dry_run']
        username = options['username']
        benefit_plan_id = options['benefit_plan_id']
        payroll_id = options['payroll_id']
        benefit_type = options['benefit_type']

        self.stdout.write(f"[import] Ficheiro : {excel_path}")
        self.stdout.write(f"[import] dry-run  : {dry_run}")
        self.stdout.write(f"[import] username  : {username}")
        self.stdout.write(f"[import] Fase 2 (Beneficiary)       : {'SIM — ' + benefit_plan_id if benefit_plan_id else 'NÃO'}")
        self.stdout.write(f"[import] Fase 3 (BenefitConsumption): {'SIM — ' + payroll_id if payroll_id else 'NÃO'}")
        self.stdout.write("")

        # Carregar Payroll e BenefitPlan se necessário
        payroll = None
        benefit_plan = None

        if payroll_id:
            try:
                payroll = Payroll.objects.get(id=payroll_id, is_deleted=False)
                self.stdout.write(f"[import] Payroll encontrado: {payroll.name}")
            except Payroll.DoesNotExist:
                raise CommandError(f"Payroll não encontrado: {payroll_id}")

        if benefit_plan_id:
            try:
                from social_protection.models import BenefitPlan
                benefit_plan = BenefitPlan.objects.get(id=benefit_plan_id, is_deleted=False)
                self.stdout.write(f"[import] BenefitPlan encontrado: {benefit_plan.name}")
            except ImportError:
                raise CommandError("Módulo social_protection não disponível")
            except Exception:
                raise CommandError(f"BenefitPlan não encontrado: {benefit_plan_id}")

        # Pré-carregar localidades
        location_cache = self._build_location_cache()

        # Ler Excel
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        sheet_name = options['sheet'] or wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            raise CommandError("Excel está vazio")

        headers = [str(h).strip() if h is not None else '' for h in rows[0]]
        self.stdout.write(f"[import] Colunas detectadas: {headers}")
        self.stdout.write(f"[import] Total de linhas de dados: {len(rows) - 1}")
        self.stdout.write("")

        def col(row, name, default=None):
            try:
                idx = headers.index(name)
                v = row[idx]
                return v if v is not None else default
            except ValueError:
                return default

        # Contadores
        stats = {
            'total': 0,
            'ind_created': 0, 'ind_updated': 0,
            'group_created': 0, 'group_updated': 0,
            'gi_created': 0,
            'benef_created': 0,
            'bc_created': 0,
            'errors': [],
        }

        for line_num, row in enumerate(rows[1:], start=2):
            if all(v is None for v in row):
                continue

            stats['total'] += 1

            nome_chefe = str(col(row, 'nome_chefe', '')).strip()
            distrito = str(col(row, 'distrito', '')).strip().upper()
            agregado_id = str(col(row, 'agregado_id', '')).strip()
            sexo = str(col(row, 'sexo_chefe', '')).strip().upper()
            data_n_raw = col(row, 'data_n')
            bi_raw = col(row, 'bilhete_didentidade')
            pan = str(col(row, 'pan', '')).strip()
            subconta = str(col(row, 'subconta', '')).strip()
            nib = str(col(row, 'NIBs', '')).strip()
            valor_raw = col(row, 'valor')
            categoria = str(col(row, 'categoria_eligibilidade', '')).strip()
            index_csu = str(col(row, 'INDEX CSU2026', '')).strip()

            if not nome_chefe or not agregado_id:
                stats['errors'].append(f"Linha {line_num}: nome ou agregado_id em falta — ignorado")
                continue

            # Processar campos
            first_name, last_name = self._split_name(nome_chefe)
            dob = self._parse_date(data_n_raw)
            bi_str = self._normalise_bi(bi_raw)
            amount = self._parse_decimal(valor_raw)
            location = location_cache.get(distrito)

            try:
                with transaction.atomic():
                    # ──────────────────────────────
                    # FASE 1A — Individual
                    # ──────────────────────────────
                    individual, ind_created = self._get_or_create_individual(
                        agregado_id=agregado_id,
                        bi_str=bi_str,
                        first_name=first_name,
                        last_name=last_name,
                        dob=dob,
                        nib=nib,
                        sexo=sexo,
                        pan=pan,
                        subconta=subconta,
                        categoria=categoria,
                        index_csu=index_csu,
                        location=location,
                        username=username,
                        dry_run=dry_run,
                    )

                    if ind_created:
                        stats['ind_created'] += 1
                        action = 'CRIADO'
                    else:
                        stats['ind_updated'] += 1
                        action = 'ACTUALIZADO'

                    self.stdout.write(
                        f"  {'[DRY] ' if dry_run else ''}Linha {line_num}: {nome_chefe} "
                        f"(agregado={agregado_id}) → Individual {action}"
                        + (f" | NIB={nib[:6]}***" if nib else " | SEM NIB")
                    )

                    # ──────────────────────────────
                    # FASE 1B — Group
                    # ──────────────────────────────
                    group, grp_created = self._get_or_create_group(
                        agregado_id=agregado_id,
                        location=location,
                        username=username,
                        dry_run=dry_run,
                    )
                    if grp_created:
                        stats['group_created'] += 1

                    # ──────────────────────────────
                    # FASE 1C — GroupIndividual
                    # ──────────────────────────────
                    gi_created = self._ensure_group_individual(
                        group=group,
                        individual=individual,
                        username=username,
                        dry_run=dry_run,
                    )
                    if gi_created:
                        stats['gi_created'] += 1

                    # ──────────────────────────────
                    # FASE 2 — Beneficiary
                    # ──────────────────────────────
                    if benefit_plan and not dry_run:
                        benef_created = self._ensure_beneficiary(
                            individual=individual,
                            benefit_plan=benefit_plan,
                            username=username,
                        )
                        if benef_created:
                            stats['benef_created'] += 1

                    # ──────────────────────────────
                    # FASE 3 — BenefitConsumption
                    # ──────────────────────────────
                    if payroll and amount is not None and not dry_run:
                        bc_created = self._ensure_benefit_consumption(
                            individual=individual,
                            payroll=payroll,
                            agregado_id=agregado_id,
                            amount=amount,
                            benefit_type=benefit_type,
                            username=username,
                        )
                        if bc_created:
                            stats['bc_created'] += 1

                    if dry_run:
                        raise _DryRunRollback()

            except _DryRunRollback:
                pass
            except Exception as exc:
                msg = f"Linha {line_num}: {nome_chefe} — ERRO: {exc}"
                stats['errors'].append(msg)
                self.stdout.write(self.style.ERROR(f"  {msg}"))
                logger.exception("[import_beneficiarios_excel] %s", msg)

        # ── Resumo ──
        self.stdout.write("")
        self.stdout.write("=" * 65)
        self.stdout.write(f"Total processado          : {stats['total']}")
        self.stdout.write(f"Indivíduos criados        : {stats['ind_created']}")
        self.stdout.write(f"Indivíduos actualizados   : {stats['ind_updated']}")
        self.stdout.write(f"Grupos criados            : {stats['group_created']}")
        self.stdout.write(f"Ligações grupo criadas    : {stats['gi_created']}")
        if benefit_plan_id:
            self.stdout.write(f"Beneficiários criados     : {stats['benef_created']}")
        if payroll_id:
            self.stdout.write(f"BenefitConsumption criados: {stats['bc_created']}")
        if dry_run:
            self.stdout.write(self.style.WARNING("MODO DRY-RUN — nenhuma alteração foi guardada"))
        if stats['errors']:
            self.stdout.write(self.style.ERROR(f"\nERROS ({len(stats['errors'])}):"))
            for e in stats['errors']:
                self.stdout.write(self.style.ERROR(f"  {e}"))

    # ────────────────────────────────────────────────────────
    # Helpers
    # ────────────────────────────────────────────────────────

    def _build_location_cache(self):
        from location.models import Location
        cache = {}
        for loc in Location.objects.filter(validity_to__isnull=True):
            name_upper = loc.name.upper() if loc.name else ''
            cache[name_upper] = loc
            for canonical, aliases in DISTRICT_ALIASES.items():
                if name_upper in aliases:
                    cache[canonical] = loc
        return cache

    def _split_name(self, full_name):
        parts = full_name.strip().split(' ', 1)
        first_name = parts[0] if parts else ''
        last_name = parts[1] if len(parts) > 1 else ''
        return first_name, last_name

    def _parse_date(self, raw):
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        if isinstance(raw, str):
            for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                try:
                    return datetime.strptime(raw.strip(), fmt).date()
                except ValueError:
                    pass
        return None

    def _normalise_bi(self, raw):
        if raw is None:
            return None
        if isinstance(raw, float):
            return str(int(raw))
        return str(raw).strip()

    def _parse_decimal(self, raw):
        if raw is None:
            return None
        try:
            return Decimal(str(raw))
        except InvalidOperation:
            return None

    def _get_or_create_individual(self, *, agregado_id, bi_str, first_name, last_name,
                                   dob, nib, sexo, pan, subconta, categoria, index_csu,
                                   location, username, dry_run):
        from individual.models import Individual

        individual = None
        created = False

        # Match 1: json_ext.agregado_id (Individual não tem campo 'code')
        if agregado_id:
            individual = Individual.objects.filter(
                json_ext__agregado_id=agregado_id, is_deleted=False
            ).first()

        # Match 2: json_ext BI
        if individual is None and bi_str:
            individual = (
                Individual.objects.filter(
                    json_ext__bilhete_didentidade=bi_str, is_deleted=False
                ).first()
                or Individual.objects.filter(
                    json_ext__bilhete_didentidade=int(bi_str) if bi_str.isdigit() else None,
                    is_deleted=False
                ).first()
            )

        if individual is None:
            individual = Individual(
                first_name=first_name,
                last_name=last_name,
            )
            created = True

        # Actualizar campos
        individual.first_name = first_name
        individual.last_name = last_name
        if dob:
            individual.dob = dob
        elif created and not individual.dob:
            # dob é obrigatório no modelo — usar data placeholder quando ausente no Excel
            individual.dob = date(1900, 1, 1)
        if nib:
            individual.nib = nib
        if location:
            individual.location = location

        json_ext = individual.json_ext or {}
        if bi_str:
            json_ext['bilhete_didentidade'] = bi_str
        if pan:
            json_ext['pan'] = pan
        if subconta:
            json_ext['subconta'] = subconta
        if sexo:
            json_ext['sexo'] = sexo
        if categoria:
            json_ext['categoria_eligibilidade'] = categoria
        if index_csu:
            json_ext['index_csu2026'] = index_csu
        json_ext['agregado_id'] = agregado_id
        individual.json_ext = json_ext

        if not dry_run:
            # HistoryModel.save() recusa guardar se nada mudou — verificar is_dirty()
            if created or individual.is_dirty():
                individual.save(username=username)

        return individual, created

    def _get_or_create_group(self, *, agregado_id, location, username, dry_run):
        from individual.models import Group

        group = Group.objects.filter(code=agregado_id, is_deleted=False).first()
        created = False

        if group is None:
            group = Group(code=agregado_id)
            if location:
                group.location = location
            created = True
            if not dry_run:
                group.save(username=username)
        elif location and group.location_id != location.id:
            group.location = location
            if not dry_run:
                group.save(username=username)

        return group, created

    def _ensure_group_individual(self, *, group, individual, username, dry_run):
        from individual.models import GroupIndividual

        existing = GroupIndividual.objects.filter(
            group=group, individual=individual, is_deleted=False
        ).first()

        if existing:
            return False

        if not dry_run:
            gi = GroupIndividual(
                group=group,
                individual=individual,
                role=GroupIndividual.Role.HEAD,
                recipient_type=GroupIndividual.RecipientType.PRIMARY,
            )
            gi.save(username=username)

        return True

    def _ensure_beneficiary(self, *, individual, benefit_plan, username):
        try:
            from social_protection.models import Beneficiary, BeneficiaryStatus
        except ImportError:
            logger.warning("[import] social_protection não disponível — Fase 2 ignorada")
            return False

        existing = Beneficiary.objects.filter(
            individual=individual,
            benefit_plan=benefit_plan,
            is_deleted=False,
        ).first()

        if existing:
            if existing.status != BeneficiaryStatus.ACTIVE:
                existing.status = BeneficiaryStatus.ACTIVE
                existing.save(username=username)
            return False

        benef = Beneficiary(
            individual=individual,
            benefit_plan=benefit_plan,
            status=BeneficiaryStatus.ACTIVE,
        )
        benef.save(username=username)
        return True

    def _ensure_benefit_consumption(self, *, individual, payroll, agregado_id,
                                     amount, benefit_type, username):
        from payroll.models import BenefitConsumption, BenefitConsumptionStatus, PayrollBenefitConsumption

        # Verificar se já existe para este payroll e individual
        existing = BenefitConsumption.objects.filter(
            individual=individual,
            payrollbenefitconsumption__payroll=payroll,
            is_deleted=False,
        ).first()

        if existing:
            return False

        today = date.today()
        code = f"{agregado_id}-{today.strftime('%Y%m')}"

        bc = BenefitConsumption(
            individual=individual,
            code=code,
            amount=amount,
            type=benefit_type,
            status=BenefitConsumptionStatus.ACCEPTED,
            date_due=today,
        )
        bc.save(username=username)

        pbc = PayrollBenefitConsumption(
            payroll=payroll,
            benefit=bc,
        )
        pbc.save(username=username)

        return True


class _DryRunRollback(Exception):
    pass
